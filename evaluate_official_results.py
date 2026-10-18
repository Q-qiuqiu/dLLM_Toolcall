"""CPU-only official scoring of saved dLLM/AR outputs; never rerun generation."""

from __future__ import annotations

import argparse
from collections import Counter
from copy import deepcopy
import csv
from datetime import datetime
from importlib.metadata import version
import json
import logging
from pathlib import Path
import sys

from official_eval_adapters import (
    ROOT, adapt_calls, check_bfcl, enable_local_dependencies, load_taskbench,
    node_projection, resource_preflight, sha256, taskbench_prediction, verify_sources,
)

enable_local_dependencies()

from json_repair import repair_json
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from bfcl_data import iter_samples
from summarize_experiment import load_records
from taskbench_data import iter_taskbench_samples


MODES = ('original', 'json_repair')
METHODS = ('LLaDA Natural', 'OMTP + Fallback', 'Dual OMTP + Fallback')
AR_METHOD = 'AR Natural (Llama 3 8B)'
DOMAINS = ('bfcl', 'dailylifeapis', 'huggingface', 'multimedia')
SUMMARY_FIELDS = [
    'dataset', 'method', 'cache_mode', 'output_mode', 'group', 'sample_count',
    'native_evaluable_gt_count', 'gt_annotation_error_count',
    'json_parseable_count', 'structural_adapter_failure_count',
    'resource_preflight_failure_count', 'unknown_tool_name_query_count',
    'bfcl_ast_correct_count', 'bfcl_ast_accuracy',
    'n_precision', 'n_recall', 'n_f1', 'e_f1', 't_f1', 'v_f1',
    'n_f1_name_projection', 'n_f1_full_adapter', 'name_projection_support', 'normalized_edit_distance',
    'official_node_support', 'official_link_support', 'official_argument_support',
    'runtime_failure_count', 'ar_context_limit_count', 'ar_eos_completion_count',
    'official_score_file',
]


def save_json(path, value):
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open('x', encoding='utf-8') as handle:
        json.dump(value, handle, ensure_ascii=False, indent=2, allow_nan=False)
        handle.write('\n')


def save_jsonl(path, records):
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open('x', encoding='utf-8') as handle:
        for record in records:
            handle.write(json.dumps(record, ensure_ascii=False, allow_nan=False) + '\n')


def run_paths(results_root, domain, *, include_ar=False):
    if domain == 'bfcl':
        names = ['bfcl_llada15_natural_only_full200',
                 'bfcl_llada15_prefixfix_nogate_full200_20261016',
                 'bfcl_llada15_dual_omtp_full200']
    else:
        prefix = f'taskbench_{domain}_llada15_'
        names = [prefix + 'natural_only_formal200',
                 prefix + 'prefixfix_gated_formal200_20261016',
                 prefix + 'dual_omtp_formal200']
    paths = [(method, results_root / name / 'samples.jsonl') for method, name in zip(METHODS, names)]
    if include_ar:
        ar_name = ('bfcl_llama3_ar_natural_eos_full200' if domain == 'bfcl' else
                   f'taskbench_{domain}_llama3_ar_natural_eos_formal200')
        paths.append((AR_METHOD, results_root / ar_name / 'samples.jsonl'))
    return paths


def inference_cache_label(meta):
    if meta['args'].get('inference_backend') == 'huggingface_ar':
        return 'ar_kv'
    return meta['args'].get('cache_mode', 'none')


def inputs_for(domain):
    if domain == 'bfcl':
        questions = ROOT / 'data/bfcl/BFCL_v4_parallel_multiple.json'
        answers = ROOT / 'data/bfcl/BFCL_v4_parallel_multiple_answers.json'
        return {s.sample_id: s for s in iter_samples(questions, answers)}, {
            'questions_sha256': (questions, sha256(questions)),
            'answers_sha256': (answers, sha256(answers)),
        }
    base = ROOT / 'data/taskbench' / f'data_{domain}'
    manifest, catalog = base / 'splits/formal200.jsonl', base / 'tool_desc.json'
    return {s.sample_id: s for s in iter_taskbench_samples(manifest, catalog)}, {
        'manifest_sha256': (manifest, sha256(manifest)),
        'tool_catalog_sha256': (catalog, sha256(catalog)),
    }


def original_taskbench_labels(domain, samples):
    import pyarrow.parquet as parquet
    base = ROOT / 'data/taskbench' / f'data_{domain}'
    raw = parquet.read_table(base / 'improved.parquet').to_pylist()
    wanted = {s.source_id for s in samples.values()}
    by_id = {str(r['id']): r for r in raw if str(r['id']) in wanted}
    if len(by_id) != len(wanted):
        raise ValueError('missing/duplicate original TaskBench IDs')
    labels = []
    for sample in samples.values():
        r = by_id[sample.source_id]
        nodes, links = json.loads(r['tool_nodes']), json.loads(r['tool_links'])
        if [n['task'] for n in nodes] != [c['name'] for c in sample.ground_truth_calls]:
            raise ValueError(f'original node order changed: {sample.sample_id}')
        endpoints = [{'source': str(e['source']), 'target': str(e['target'])} for e in links]
        if endpoints != sample.dependencies:
            raise ValueError(f'original dependency edges changed: {sample.sample_id}')
        labels.append({'id': sample.source_id, 'type': r['type'],
                       'task_steps': json.loads(r['tool_steps']), 'task_nodes': nodes,
                       'task_links': links})
    return labels, json.loads((base / 'tool_desc.json').read_text()), base / 'improved.parquet'


def invoke_taskbench(data_dir, stem, catalog, dependency_type, log_path, *, projection=False):
    module = load_taskbench()
    tool_map = {n['id']: i + 1 for i, n in enumerate(catalog['nodes'])}
    reverse = {i + 1: n['id'] for i, n in enumerate(catalog['nodes'])}
    reverse[0] = 'NEGATIVE'; tool_map['<PAD>'] = -1
    outputs = ({n['id']: n['output-type'][0] if n['output-type'] else 'none' for n in catalog['nodes']}
               if dependency_type == 'resource' else None)
    handler = logging.FileHandler(log_path, encoding='utf-8')
    module.logger.addHandler(handler)
    results = {}
    try:
        for group in ('overall', 'chain', 'dag'):
            module.evaluate(str(data_dir), 'predictions', stem, group, 'overall',
                            ['f1'] if projection else ['f1', 'ed', 'link', 'argument'],
                            catalog, tool_map, outputs, reverse, results, dependency_type, None)
    finally:
        module.logger.removeHandler(handler)
        handler.close()
    for group in ('overall', 'chain', 'dag'):
        score = results[group + '_overall']
        for key in ('node_supports', 'link_supports', 'argument_supports'):
            if score[key] != score['all_samples']:
                raise RuntimeError(f'upstream evaluator silently excluded samples: {stem}/{group}/{key}; see {log_path}')
    return results


def sanity_taskbench(data_dir, labels, catalog, dependency_type):
    if dependency_type == 'resource':
        for label in labels:
            if error := resource_preflight(label['task_nodes']):
                raise ValueError(f'original TaskBench GT has unsupported resource argument: {label["id"]}: {error}')
    for stem, empty in [('sanity_perfect_gt_NOT_A_MODEL', False), ('sanity_empty_NOT_A_MODEL', True)]:
        predictions = [{'id': r['id'], 'result': {'task_nodes': [] if empty else deepcopy(r['task_nodes']),
                                                'task_links': [] if empty else deepcopy(r['task_links'])}}
                       for r in labels]
        save_jsonl(data_dir / 'predictions' / (stem + '.json'), predictions)
        scores = invoke_taskbench(data_dir, stem, catalog, dependency_type, data_dir / (stem + '.log'))
        save_json(data_dir / (stem + '_scores.json'), scores)
        for group in scores.values():
            for key in ['node_micro_f1_no_matching', 'link_binary_f1',
                        'argument_task_argname_binary_f1_no_matching',
                        'argument_task_argname_value_binary_f1_no_matching']:
                expected = 0.0 if empty else 1.0
                if abs(group[key] - expected) > 1e-12:
                    raise AssertionError(f'official evaluator sanity failure: {dependency_type}/{stem}/{key}={group[key]}')


def ground_truth_error(label, dependency_type):
    for index, node in enumerate(label['task_nodes']):
        if not isinstance(node.get('arguments'), list):
            return f'node_{index}:missing_or_invalid_arguments_list'
    if dependency_type == 'resource':
        return resource_preflight(label['task_nodes'])
    bad = []
    for i, node in enumerate(label['task_nodes']):
        for j, argument in enumerate(node['arguments']):
            if not isinstance(argument, dict) or 'name' not in argument or 'value' not in argument:
                bad.append({'node_index': i, 'argument_index': j, 'argument': argument})
    return 'missing_named_argument_value:' + json.dumps(bad, ensure_ascii=False) if bad else None


def mode_text(row, mode):
    text = row.get('output_before_eos', row['raw_output'])
    if mode == 'json_repair':
        try:
            json.loads(text)
        except (ValueError, TypeError):
            text = str(repair_json(text, ensure_ascii=False))
    return text


def evaluate_run(domain, method, path, samples, checksums, output, taskbench_context=None):
    meta, records = load_records(path)
    if {r['sample_id'] for r in records} != set(samples) or len(records) != len(samples):
        raise ValueError(f'not the full fixed sample list: {path}')
    for key, (_, checksum) in checksums.items():
        if meta[key] != checksum:
            raise ValueError(f'data changed since generation: {path}/{key}')
    schema = meta['args'].get('output_schema', 'standard')
    cache = inference_cache_label(meta)
    if method == AR_METHOD and meta['args'].get('inference_backend') != 'huggingface_ar':
        raise ValueError(f'AR input is not a HuggingFace AR run: {path}')
    summaries, audit = [], []
    stem_base = path.parent.name
    for mode in MODES:
        stem = stem_base + '__' + mode
        full_predictions, projections, mode_records = [], [], []
        for row in records:
            sample = samples[row['sample_id']]
            if row['candidate_tools'] != sample.candidate_names:
                raise ValueError(f'candidate order/content changed: {row["sample_id"]}')
            if row['ground_truth_calls'] != sample.ground_truth_calls:
                raise ValueError(f'GT changed since generation: {row["sample_id"]}')
            repair_error = None
            try:
                text = mode_text(row, mode)
            except Exception as error:
                text = ''; repair_error = f'{type(error).__name__}: {error}'
            adapted = adapt_calls(text, schema, sample.functions)
            try:
                json.loads(text)
                syntax = True
            except (ValueError, TypeError):
                syntax = False
            item = {'dataset': domain, 'method': method, 'cache_mode': cache, 'output_mode': mode,
                    'sample_id': row['sample_id'], 'graph_type': row.get('graph_type'),
                    'json_parseable': syntax, 'adapter_error': adapted.error,
                    'repair_error': repair_error, 'canonical_calls': adapted.calls,
                    'runtime_error': row.get('error'),
                    'generation_stop_reason': row.get('stop_reason'),
                    'natural_generation_complete': row.get('natural_generation_complete'),
                    'unknown_tool_names': sorted({c['name'] for c in adapted.calls} - set(sample.candidate_names)),
                    'resource_preflight_error': None, 'original_result_file': str(path.resolve())}
            if domain == 'bfcl':
                result = check_bfcl(sample, adapted)
                item['official_ast_result'] = result
            else:
                data_dir, labels, catalog, dep_type, gt_errors = taskbench_context
                native = taskbench_prediction(adapted.calls, dep_type)
                projections.append({'id': sample.source_id, 'result': node_projection(native)})
                if dep_type == 'resource':
                    item['resource_preflight_error'] = resource_preflight(native['task_nodes'])
                if item['resource_preflight_error']:
                    native = {'task_nodes': [], 'task_links': []}
                full_predictions.append({'id': sample.source_id, 'result': native})
                item['official_prediction'] = native
                item['gt_annotation_error'] = gt_errors.get(sample.source_id)
            mode_records.append(item)
        audit.extend(mode_records)
        if domain == 'bfcl':
            count = sum(r['official_ast_result']['valid'] for r in mode_records)
            destination = output / 'bfcl' / (stem + '_scores.json')
            save_json(destination, {'accuracy': count / len(records), 'correct_count': count,
                                    'sample_count': len(records),
                                    'errors': dict(Counter(r['official_ast_result'].get('error_type', 'valid') for r in mode_records))})
            groups = [('overall', mode_records, {'bfcl_ast_correct_count': count, 'bfcl_ast_accuracy': count / len(records),
                                                 'native_evaluable_gt_count': len(records), 'gt_annotation_error_count': 0})]
        else:
            data_dir, labels, catalog, dep_type, gt_errors = taskbench_context
            native_dir, names_dir = data_dir / 'native', data_dir / 'names_only'
            save_jsonl(native_dir / 'predictions' / (stem + '.json'), full_predictions)
            projection_stem = stem + '__names_only'
            save_jsonl(names_dir / 'predictions' / (projection_stem + '.json'), projections)
            scores = invoke_taskbench(native_dir, stem, catalog, dep_type, data_dir / (stem + '.log'))
            name_scores = invoke_taskbench(names_dir, projection_stem, catalog, dep_type,
                                          data_dir / (projection_stem + '.log'), projection=True)
            destination = data_dir / (stem + '_scores.json')
            save_json(destination, scores)
            save_json(data_dir / (projection_stem + '_scores.json'), name_scores)
            groups = []
            for group in ['overall', 'chain', 'dag']:
                s = scores[group + '_overall']
                members = [r for r in mode_records if group == 'overall' or r['graph_type'] == group]
                gt_bad_count = sum(r.get('gt_annotation_error') is not None for r in members)
                ns = name_scores[group + '_overall']
                if s['all_samples'] != len(members) - gt_bad_count or ns['all_samples'] != len(members):
                    raise AssertionError(f'wrong official denominator: {stem}/{group}')
                groups.append((group, members, {
                    'native_evaluable_gt_count': s['all_samples'], 'gt_annotation_error_count': gt_bad_count,
                    'n_precision': ns['node_micro_precision_no_matching'],
                    'n_recall': ns['node_micro_recall_no_matching'], 'n_f1': ns['node_micro_f1_no_matching'],
                    'e_f1': s['link_binary_f1'], 't_f1': s['argument_task_argname_binary_f1_no_matching'],
                    'v_f1': s['argument_task_argname_value_binary_f1_no_matching'],
                    'n_f1_name_projection': ns['node_micro_f1_no_matching'],
                    'n_f1_full_adapter': s['node_micro_f1_no_matching'], 'name_projection_support': ns['node_supports'],
                    'normalized_edit_distance': s['edit_distance'], 'official_node_support': s['node_supports'],
                    'official_link_support': s['link_supports'], 'official_argument_support': s['argument_supports'],
                }))
        for group, members, scores_row in groups:
            summaries.append(dict(dataset=domain, method=method, cache_mode=cache, output_mode=mode,
                                  group=group, sample_count=len(members),
                                  json_parseable_count=sum(r['json_parseable'] for r in members),
                                  structural_adapter_failure_count=sum(r['adapter_error'] is not None for r in members),
                                  resource_preflight_failure_count=sum(r['resource_preflight_error'] is not None for r in members),
                                  unknown_tool_name_query_count=sum(bool(r['unknown_tool_names']) for r in members),
                                  runtime_failure_count=sum(bool(r['runtime_error']) for r in members),
                                  ar_context_limit_count=(sum(r['generation_stop_reason'] == 'context_limit' for r in members)
                                                          if cache == 'ar_kv' else None),
                                  ar_eos_completion_count=(sum(r['generation_stop_reason'] == 'eos' for r in members)
                                                           if cache == 'ar_kv' else None),
                                  official_score_file=str(destination.resolve()), **scores_row))
        print(json.dumps({k: summaries[-len(groups)].get(k) for k in [
            'dataset', 'method', 'output_mode', 'sample_count', 'bfcl_ast_accuracy',
            'n_f1', 'e_f1', 't_f1', 'v_f1', 'resource_preflight_failure_count']}, ensure_ascii=False), flush=True)
    return summaries, audit


def export_workbook(path, summaries, parity, metadata):
    book = Workbook()
    readme = book.active; readme.title = '阅读说明'
    notes = [
        ('范围', '保存的LLaDA-1.5结果；No Cache Natural、No Cache OMTP、Dual OMTP。启用--include-ar时另含Llama 3 8B AR Natural。没有重新生成、启动模型或改动在线算法/时间。'),
        ('分数性质', '真实调用固定版本官方checker/evaluator，但使用自定义JSON输出适配和固定测试子集；不是官方榜单原生Prompt/全量分数。'),
        ('Original', 'output_before_eos（若有）否则raw_output，严格JSON解析；不能适配的请求作为空预测计失败，不从分母删除。'),
        ('Repair', '仅对非法JSON离线运行一次json-repair，再执行相同适配及官方评分；原始输出/在线事件/时间不变。'),
        ('AR公平性', 'AR复用完全相同的评分源码、JSON适配、Repair和GT排除清单。使用已完成的自然EOS实验，不补加128-token预算；上下文上限失败不删除、不当作自然结束，次数单列。ar_kv是AR原生KV Cache，不是dLLM Dual Cache。'),
        ('BFCL AST', '直接调用未修改的官方ast_checker，Python parallel_multiple，无序且重复调用敏感。模型注册桥仅设置underscore_to_dot=False，不加载API provider。'),
        ('TaskBench参数适配', 'DailyLife还原named arguments；HuggingFace/Multimedia还原input_N的原始位置参数。保留未知工具、额外参数、错误类型和缺失参数，由官方评分。'),
        ('TaskBench依赖适配', 'Resource领域由官方脚本根据预测参数引用生成边；DailyLife没有显式task_links输出，只能从已生成的精确<node-N>数据流引用恢复边，不根据GT补时序依赖。'),
        ('官方脚本缺陷处理', 'Resource原脚本except/pass可能将上一条甚至GT参数复用于错误引用。未改官方代码；运行前验证参数，非法引用/不支持类型整条预测按空输出计失败。原始预测与原因保存在evaluation_samples.jsonl。'),
        ('n_f1 / n_f1_name_projection', '对全部200条只投影工具节点，GT和预测都移除参数，再调用同一官方脚本计算节点n-F1；避免参数坏标注或解析异常影响工具名称指标。'),
        ('n_f1_full_adapter', '附加：原生完整参数适配路径中的n-F1，分母为GT原生可评子集；非法资源引用的预测为空，因此它还受到参数适配影响。不要与主n_f1混称。'),
        ('GT坏标注', 'DailyLife 1条、HuggingFace 2条、Multimedia 3条不满足原官方参数解析要求。节点n-F1仍为200条；原生e/t/v-F1分别为199/198/197条，单独保存坏标注及原因，绝不补GT。'),
        ('全样本审计', '模型JSON/适配失败作为空预测保留；节点路径检查各200条，原生路径检查有效GT子集；任何未声明的静默跳过都会报错停止。'),
        ('Sanity', '每个TaskBench领域分别用GT预测/空预测检查四种官方F1为1/0，明确非模型结果。'),
        ('未知工具', '官方TaskBench n-F1只计目录内正类，未知名称按0类不贡献节点FP；未知名称请求数另列，不能只看n-F1忽略非法工具。'),
        ('Resource参数语义', 'HuggingFace/Multimedia的官方t-F1比较工具+输入资源类型；v-F1进一步比较资源值。官方根据文件扩展名和来源工具output-type推断类型，并将<node-N>解析为来源工具名，不是对input_N字段名或引用编号做字面AST匹配。'),
        ('NED', 'normalized_edit_distance对应官方edit_distance=1−Levenshtein.ratio；越低越好。Chain更有解释性，DAG顺序不唯一。'),
        ('未评指标', '未生成任务分解文本，因此不评ROUGE/BERTScore；未执行工具，因此不报告真实任务执行成功率。'),
        ('官方版本', json.dumps(metadata['official_sources'], ensure_ascii=False)),
    ]
    readme.append(['项目', '说明'])
    for r in notes: readme.append(r)
    for name, rows in [('官方指标_全部分组', summaries), ('官方指标_Overall', [r for r in summaries if r['group'] == 'overall'])]:
        ws = book.create_sheet(name); ws.append(SUMMARY_FIELDS)
        for row in rows: ws.append([row.get(k) for k in SUMMARY_FIELDS])
    ws = book.create_sheet('输出一致性')
    ws.append(['dataset', 'sample_count', 'nocache_natural_omtp_raw_identical_count', 'official_scores_identical'])
    for row in parity:
        ws.append([row[k] for k in ['dataset', 'sample_count', 'nocache_natural_omtp_raw_identical_count', 'official_scores_identical']])
    for ws in book:
        ws.freeze_panes = 'C2' if ws.max_column > 2 else 'A2'; ws.auto_filter.ref = ws.dimensions
        ws.sheet_view.showGridLines = False; ws.row_dimensions[1].height = 38
        for col in range(1, ws.max_column + 1):
            h = str(ws.cell(1, col).value or '')
            ws.column_dimensions[get_column_letter(col)].width = 28 if ws.max_column > 2 else (26 if col == 1 else 110)
            if 'file' in h: ws.column_dimensions[get_column_letter(col)].width = 70
            for row in range(1, ws.max_row + 1):
                c = ws.cell(row, col); c.font = Font(name='等线', size=10, bold=row == 1, color='FFFFFF' if row == 1 else '000000')
                c.alignment = Alignment(wrap_text=True, vertical='top')
                if row == 1: c.fill = PatternFill('solid', fgColor='17365D')
                elif row % 2 == 0: c.fill = PatternFill('solid', fgColor='F0F5FA')
                if isinstance(c.value, str): c.data_type = 's'
                if row > 1 and isinstance(c.value, float):
                    c.number_format = '0.00%' if h.endswith(('_accuracy', '_precision', '_recall', '_f1', '_projection')) else '0.0000'
    book.save(path)


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--results-root', type=Path, default=ROOT / 'results')
    parser.add_argument('--output-dir', type=Path, required=True)
    parser.add_argument('--domains', nargs='+', choices=DOMAINS, default=DOMAINS)
    parser.add_argument('--include-ar', action='store_true', help='Also score the four saved Llama 3 AR Natural EOS runs')
    args = parser.parse_args()
    output = args.output_dir.resolve()
    if output.exists(): raise FileExistsError(f'refusing to overwrite {output}')
    output.mkdir(parents=True)
    source_manifest = verify_sources()
    metadata = {'created_at': datetime.now().astimezone().isoformat(), 'command': sys.argv,
                'python': sys.version, 'official_sources': source_manifest,
                'dependencies': {p: version(p) for p in ['numpy', 'scipy', 'scikit-learn', 'pyarrow',
                                                        'datasets', 'Levenshtein', 'json-repair', 'openpyxl']},
                'adapter_sha256': sha256(ROOT / 'official_eval_adapters.py'),
                'runner_sha256': sha256(Path(__file__)), 'inputs': [],
                'no_inference': True, 'raw_and_online_records_unchanged': True, 'include_ar': args.include_ar,
                'bfcl_model_registry_bridge': {'model_name': 'llada-offline', 'underscore_to_dot': False}}
    all_summaries, all_audits, parity = [], [], []
    for domain in args.domains:
        samples, checksums = inputs_for(domain)
        context = None
        if domain != 'bfcl':
            labels, catalog, parquet = original_taskbench_labels(domain, samples)
            data_dir = output / ('taskbench_' + domain)
            dep_type = 'temporal' if domain == 'dailylifeapis' else 'resource'
            gt_errors = {r['id']: error for r in labels if (error := ground_truth_error(r, dep_type))}
            save_jsonl(data_dir / 'original_source_data.json', labels)
            save_json(data_dir / 'gt_annotation_errors.json', [
                {'id': r['id'], 'error': gt_errors[r['id']], 'original_label': r}
                for r in labels if r['id'] in gt_errors])
            native_dir, names_dir = data_dir / 'native', data_dir / 'names_only'
            valid_labels = [r for r in labels if r['id'] not in gt_errors]
            save_jsonl(native_dir / 'data.json', valid_labels)
            save_json(native_dir / 'tool_desc.json', catalog)
            projected_labels = [dict(r, task_nodes=node_projection(r)['task_nodes']) for r in labels]
            save_jsonl(names_dir / 'data.json', projected_labels)
            save_json(names_dir / 'tool_desc.json', catalog)
            sanity_taskbench(native_dir, valid_labels, catalog, dep_type)
            context = (data_dir, labels, catalog, dep_type, gt_errors)
            metadata.setdefault('original_parquet', {})[domain] = {'path': str(parquet), 'sha256': sha256(parquet)}
            metadata.setdefault('gt_annotation_errors', {})[domain] = gt_errors
        by_method = {}
        for method, path in run_paths(args.results_root, domain, include_ar=args.include_ar):
            if not path.exists():
                raise FileNotFoundError(path)
            original_sha = sha256(path)
            summaries, audit = evaluate_run(domain, method, path, samples, checksums, output, context)
            if sha256(path) != original_sha: raise AssertionError('input JSONL was modified')
            metadata['inputs'].append({'dataset': domain, 'method': method, 'path': str(path.resolve()), 'sha256': original_sha})
            all_summaries.extend(summaries); all_audits.extend(audit); by_method[method] = summaries
        _, baseline = load_records(run_paths(args.results_root, domain)[0][1])
        _, omtp = load_records(run_paths(args.results_root, domain)[1][1])
        omtp_index = {r['sample_id']: r for r in omtp}
        identical = sum(r['raw_output'] == omtp_index[r['sample_id']]['raw_output'] for r in baseline)
        fields = ['output_mode', 'group', 'bfcl_ast_accuracy', 'n_f1', 'e_f1', 't_f1', 'v_f1', 'n_f1_name_projection']
        same_scores = ([tuple(r.get(k) for k in fields) for r in by_method[METHODS[0]]] ==
                       [tuple(r.get(k) for k in fields) for r in by_method[METHODS[1]]])
        if identical == len(baseline) and not same_scores:
            raise AssertionError('identical outputs received different official scores')
        parity.append({'dataset': domain, 'sample_count': len(baseline),
                       'nocache_natural_omtp_raw_identical_count': identical, 'official_scores_identical': same_scores})
    save_json(output / 'metadata.json', metadata)
    save_json(output / 'output_parity.json', parity)
    save_jsonl(output / 'evaluation_samples.jsonl', all_audits)
    with (output / 'official_summary.csv').open('x', encoding='utf-8-sig', newline='') as handle:
        writer = csv.DictWriter(handle, fieldnames=SUMMARY_FIELDS)
        writer.writeheader(); writer.writerows(all_summaries)
    export_workbook(output / 'official_metrics.xlsx', all_summaries, parity, metadata)
    print(f'Completed: {len(all_audits)} sample/mode records; results: {output}', flush=True)


if __name__ == '__main__':
    main()
