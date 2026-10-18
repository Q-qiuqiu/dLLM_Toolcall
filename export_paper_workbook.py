"""Export existing formal runs to one auditable workbook; never run inference.

Requires openpyxl, torch (existing local AST checker), and json-repair.
Official evaluator scores are deliberately NOT inferred from custom scores.
"""

from __future__ import annotations

import argparse
from collections import Counter
from datetime import datetime
import hashlib
from importlib.metadata import version
import json
from pathlib import Path
from statistics import mean

from openpyxl import Workbook, load_workbook
from openpyxl.comments import Comment
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from bfcl_data import iter_samples
from compare_ar_baseline import compare_ar
from evaluation import bfcl_compatible_ast_match
from json_repair_evaluation import evaluate_json_repair
from summarize_experiment import compare_natural_baseline, distribution, load_records
from taskbench_data import iter_taskbench_samples
from taskbench_dependency_evaluation import evaluate_semantic_dependencies
from tool_parser import extract_partial_names, parse_tool_calls


ROOT = Path(__file__).resolve().parent
METHODS = ("AR Natural", "LLaDA Natural", "OMTP + Fallback")
DATASETS = {
    "BFCL parallel_multiple": ("bfcl", {
        METHODS[0]: "bfcl_llama3_ar_natural_eos_full200",
        METHODS[1]: "bfcl_llada15_natural_only_full200",
        METHODS[2]: "bfcl_llada15_prefixfix_nogate_full200_20261016",
    }),
    "TaskBench DailyLife APIs": ("dailylifeapis", {
        METHODS[0]: "taskbench_dailylifeapis_llama3_ar_natural_eos_formal200",
        METHODS[1]: "taskbench_dailylifeapis_llada15_natural_only_formal200",
        METHODS[2]: "taskbench_dailylifeapis_llada15_prefixfix_gated_formal200_20261016",
    }),
    "TaskBench HuggingFace": ("huggingface", {
        METHODS[0]: "taskbench_huggingface_llama3_ar_natural_eos_formal200",
        METHODS[1]: "taskbench_huggingface_llada15_natural_only_formal200",
        METHODS[2]: "taskbench_huggingface_llada15_prefixfix_gated_formal200_20261016",
    }),
    "TaskBench Multimedia": ("multimedia", {
        METHODS[0]: "taskbench_multimedia_llama3_ar_natural_eos_formal200",
        METHODS[1]: "taskbench_multimedia_llada15_natural_only_formal200",
        METHODS[2]: "taskbench_multimedia_llada15_prefixfix_gated_formal200_20261016",
    }),
}
OFFICIAL_LINKS = {
    "bfcl": "https://gorilla.cs.berkeley.edu/blogs/8_berkeley_function_calling_leaderboard.html",
    "taskbench": "https://github.com/microsoft/JARVIS/blob/main/taskbench/README.md",
    "taskbench_evaluator": "https://huggingface.co/datasets/microsoft/Taskbench/blob/main/evaluate.py",
}
NA_OFFICIAL = "未运行官方 evaluator"


def quality(predicted, expected):
    """Set-level micro counts and request-level coverage, including all failures."""
    assert len(predicted) == len(expected)
    tp = sum(len(p & g) for p, g in zip(predicted, expected))
    fp = sum(len(p - g) for p, g in zip(predicted, expected))
    fn = sum(len(g - p) for p, g in zip(predicted, expected))
    n = len(predicted)
    return {
        "tp": tp, "fp": fp, "fn": fn,
        "precision": tp / (tp + fp) if tp + fp else None,
        "recall": tp / (tp + fn) if tp + fn else None,
        "f1": 2 * tp / (2 * tp + fp + fn) if 2 * tp + fp + fn else None,
        "exact_count": sum(p == g for p, g in zip(predicted, expected)),
        "exact": sum(p == g for p, g in zip(predicted, expected)) / n if n else None,
        "covered_count": sum(bool(g) and g <= p for p, g in zip(predicted, expected)),
        "nonempty_target_count": sum(bool(g) for g in expected),
        "covered": (sum(bool(g) and g <= p for p, g in zip(predicted, expected)) /
                    sum(bool(g) for g in expected)) if any(expected) else None,
        "wrong_rate": fp / (tp + fp) if tp + fp else None,
        "wrong_per_query": fp / n if n else None,
        "wrong_queries": sum(bool(p - g) for p, g in zip(predicted, expected)),
        "missing_per_query": fn / n if n else None,
    }


def event_map(row, natural=False):
    if row.get("error"):
        return {}
    if natural:
        return {tool: event["time_s"] for tool, event in row["natural_first_times"].items()}
    events = row["first_tool_events"]
    if len({e["tool"] for e in events}) != len(events):
        raise ValueError(f"duplicate first identity event: {row['sample_id']}")
    return {e["tool"]: e["time_s"] for e in events}


def complete_time(events, target):
    return max(events[t] for t in target) if target and target <= events.keys() else None


def sample_catalog(domain):
    if domain == "bfcl":
        samples = iter_samples(ROOT / "data/bfcl/BFCL_v4_parallel_multiple.json",
                               ROOT / "data/bfcl/BFCL_v4_parallel_multiple_answers.json")
    else:
        base = ROOT / "data/taskbench" / f"data_{domain}"
        samples = iter_taskbench_samples(base / "splits/formal200.jsonl", base / "tool_desc.json")
    return {sample.sample_id: sample for sample in samples}


def analyze(dataset, domain, method, path, samples):
    meta, records = load_records(path)
    if len(records) != len({r['sample_id'] for r in records}):
        raise ValueError(f"duplicate requests: {path}")
    enriched = []
    schema = meta['args'].get('output_schema', 'standard')
    for row in records:
        sample = samples[row['sample_id']]
        if set(row['candidate_tools']) != set(sample.candidate_names):
            raise ValueError(f"candidate catalog mismatch: {row['sample_id']}")
        if set(row['ground_truth_tools']) != sample.ground_truth_tool_types:
            raise ValueError(f"ground truth mismatch: {row['sample_id']}")
        online = event_map(row)
        natural = event_map(row, natural=True)
        final_names = extract_partial_names(row['raw_output'], row['candidate_tools'])
        final = set(final_names)
        gt = set(row['ground_truth_tools'])
        text = row.get('output_before_eos', row['raw_output'])
        parsed = parse_tool_calls(text, sample.candidate_names,
                                  output_schema=schema, functions=sample.functions)
        repair = evaluate_json_repair(text, sample.candidate_names, sample, output_schema=schema)
        repaired = parse_tool_calls(repair['repaired_json_text'], sample.candidate_names,
                                    output_schema=schema, functions=sample.functions)
        enrich = dict(row, online=online, natural=natural, final=final, gt=gt,
                      raw_names=final_names, strict=parsed, repaired=repaired,
                      repair=repair, gt_names=[c['name'] for c in sample.ground_truth_calls])
        for key, result in [('strict', parsed), ('repaired', repaired)]:
            calls = result.calls if result.format_valid else []
            enrich[key + '_ast'] = result.valid and bfcl_compatible_ast_match(calls, sample)
            if domain != 'bfcl':
                enrich[key + '_dependencies'] = evaluate_semantic_dependencies(
                    calls, row['dependency_edges'], sample.ground_truth_calls, output_valid=result.valid)
        enriched.append(enrich)
    return dict(dataset=dataset, domain=domain, method=method, path=path, meta=meta,
                rows=enriched, samples=samples)


def set_metrics(run, reference):
    return quality([set(r['online']) for r in run['rows']], [r[reference] for r in run['rows']])


def sheet(book, title, headers, records):
    ws = book.create_sheet(title)
    ws.append(headers)
    for record in records:
        ws.append([record.get(h) for h in headers] if isinstance(record, dict) else list(record))
    ws.freeze_panes = 'C2' if len(headers) >= 3 else 'A2'
    ws.auto_filter.ref = ws.dimensions
    ws.sheet_view.showGridLines = False
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.orientation = 'landscape'
    ws.page_setup.paperSize = ws.PAPERSIZE_A3
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.print_title_rows = '1:1'
    ws.row_dimensions[1].height = 45
    for cell in ws[1]:
        cell.font = Font(name='等线', bold=True, color='FFFFFF', size=10)
        cell.fill = PatternFill('solid', fgColor='17365D')
        cell.alignment = Alignment(vertical='center', wrap_text=True)
    for idx, header in enumerate(headers, 1):
        width = 23 if idx > 2 else 30
        if any(s in header for s in ['说明', '定义', '警告', '原因', '路径', '来源', '命令', '哈希', '参考链接']):
            width = 68
        ws.column_dimensions[get_column_letter(idx)].width = width
        is_rate = any(s in header for s in ['率', '比例', 'Precision', 'Recall', 'F1', 'Exact Match'])
        if any(s in header for s in ['分母', '数量', '个数', '次数', '样本数', '任务数', '工具数', '计数']):
            is_rate = False
        for cells in ws.iter_rows(min_row=2, min_col=idx, max_col=idx):
            cell = cells[0]
            cell.font = Font(name='等线', size=10)
            cell.alignment = Alignment(vertical='top', wrap_text=True)
            if cell.row % 2 == 0:
                cell.fill = PatternFill('solid', fgColor='F0F5FA')
            if isinstance(cell.value, str):
                # Raw/model-controlled text must never become an Excel formula.
                cell.data_type = 's'
                if cell.value.startswith('未运行') or cell.value.startswith('不可评估'):
                    cell.font = Font(name='等线', color='9C6500', size=10)
            elif isinstance(cell.value, (int, float)) and not isinstance(cell.value, bool):
                cell.number_format = '0.00%' if is_rate else ('0.000' if isinstance(cell.value, float) else '0')
        if is_rate and ws.max_row > 2:
            col = get_column_letter(idx)
            ws.conditional_formatting.add(f'{col}2:{col}{ws.max_row}', ColorScaleRule(
                start_type='num', start_value=0, start_color='FBE2D5',
                end_type='num', end_value=1, end_color='D5EAD5'))
    return ws


def add_distribution(rows, dataset, method, metric, values, eligible, reference, note=''):
    d = distribution(values)
    assert d['count'] <= eligible
    rows.append([dataset, method, metric, reference, eligible, d['count'],
                 eligible - d['count'], (eligible - d['count']) / eligible if eligible else None,
                 d['mean'], d['p50'], d['p95'], note])


def timing_rows(run):
    ds, method, rows = run['dataset'], run['method'], run['rows']
    out = []
    n = len(rows)
    add_distribution(out, ds, method, '整条请求生成墙钟时间',
                     [r.get('generation_time_s') for r in rows], n, '全请求', '包含本方法 Observer；单位秒')
    add_distribution(out, ds, method, 'Observer 处理时间/请求',
                     [r.get('observer_time_s') for r in rows], n, '全请求', '不是生成总时间差；不包含所有 hook/同步开销')
    for reference, label in [('final', '各自 Final Raw Text'), ('gt', 'Ground Truth')]:
        eligible = sum(len(r[reference]) for r in rows)
        times = [r['online'].get(t) for r in rows for t in r[reference]]
        add_distribution(out, ds, method, '逐工具类型首次识别时间', times, eligible, label,
                         '每请求×工具类型只计一次；各模型自己的 Final 不同，不可直接比较其均值')
        eligible_rows = [r for r in rows if r[reference]]
        add_distribution(out, ds, method, '完整工具类型集合覆盖时间',
                         [complete_time(r['online'], r[reference]) for r in eligible_rows],
                         len(eligible_rows), label, '允许额外错误触发；失败留空，不补0；正确性是离线标注')
    for k in range(1, 5):
        eligible_rows = [r for r in rows if len(r['gt']) >= k]
        values = []
        for r in eligible_rows:
            times = sorted(r['online'][t] for t in r['gt'] if t in r['online'])
            values.append(times[k - 1] if len(times) >= k else None)
        add_distribution(out, ds, method, f'第{k}个正确工具类型识别时间', values,
                         len(eligible_rows), 'Ground Truth', '按真实识别时间排序，不是调用数组位置；不同名称去重')
    return out


def quality_rows(run):
    out = []
    for reference, label in [('final', 'Final Raw Text（候选名称内）'), ('gt', 'Ground Truth')]:
        q = set_metrics(run, reference)
        n = len(run['rows'])
        out.append([run['dataset'], run['method'], label, n, q['tp'], q['fp'], q['fn'],
                    q['precision'], q['recall'], q['f1'], q['exact_count'], q['exact'],
                    q['covered_count'], q['nonempty_target_count'], q['covered'],
                    q['wrong_rate'], q['wrong_per_query'], q['wrong_queries'],
                    q['wrong_queries'] / n, q['missing_per_query'],
                    '不可评估：日志只保留每种工具首次事件',
                    '不可评估：日志没有完整 occurrence/位置轨迹'])
    return out


def generation_rows(run):
    out = []
    rows = run['rows']
    n = len(rows)
    raw_q = quality([r['final'] for r in rows], [r['gt'] for r in rows])
    for key, label in [('strict', 'Original / EOS前文本'), ('repaired', 'JSON Repair（离线）')]:
        parsed = [r[key] for r in rows]
        predictions = [r[key].names if r[key].format_valid else [] for r in rows]
        gold = [r['gt_names'] for r in rows]
        q = quality([set(p) for p in predictions], [set(g) for g in gold])
        tp = sum(sum((Counter(p) & Counter(g)).values()) for p, g in zip(predictions, gold))
        np = sum(map(len, predictions)); ng = sum(map(len, gold))
        exact = sum(r[key].valid and set(p) == r['gt'] for r, p in zip(rows, predictions))
        position_correct = sum(sum(pn == gn for pn, gn in zip(p, g)) for p, g in zip(predictions, gold))
        edge_f1 = None
        if run['domain'] != 'bfcl':
            deps = [r[key + '_dependencies'] for r in rows]
            etp = sum(d['edge_true_positive_count'] for d in deps)
            efp = sum(d['edge_false_positive_count'] for d in deps)
            efn = sum(d['edge_false_negative_count'] for d in deps)
            edge_f1 = 2 * etp / (2 * etp + efp + efn) if 2 * etp + efp + efn else None
        raw_syntax = []
        for r in rows:
            try:
                json.loads(r['raw_output'])
                raw_syntax.append(True)
            except (TypeError, ValueError):
                raw_syntax.append(False)
        out.append([run['dataset'], run['method'], label, n, mean(raw_syntax),
                    mean(p.json_valid for p in parsed), mean(p.format_valid for p in parsed),
                    mean(p.valid for p in parsed), q['precision'], q['recall'], q['f1'], exact / n,
                    raw_q['exact'], tp / np if np else None, tp / ng if ng else None,
                    2 * tp / (np + ng) if np + ng else None,
                    sum(r[key].valid and len(p) == len(g) for r, p, g in zip(rows, predictions, gold)) / n,
                    sum(r[key].valid and Counter(p) == Counter(g) for r, p, g in zip(rows, predictions, gold)) / n,
                    position_correct / ng if ng else None,
                    mean(r[key + '_ast'] for r in rows), edge_f1,
                    NA_OFFICIAL,
                    '所有请求为分母；调用结构失败按空调用计；未知工具保留为FP。'
                    '本地AST/语义边不等于官方分数；BFCL并行调用顺序不作官方正确性要求。'])
    return out


def paired_rows(dataset, name_a, name_b, pairs, target, metric, note):
    eligible = len(pairs)
    states = Counter({'both': 0, 'a': 0, 'b': 0, 'neither': 0})
    aa, bb, lead = [], [], []
    for a, b in pairs:
        state = 'both' if a is not None and b is not None else 'a' if a is not None else 'b' if b is not None else 'neither'
        states[state] += 1
        if state == 'both':
            aa.append(a); bb.append(b); lead.append(b - a)
    da, db, dl = (distribution(v) for v in [aa, bb, lead])
    return [dataset, name_a, name_b, target, metric, eligible, states['both'], states['a'],
            states['b'], states['neither'], da['mean'], da['p50'], da['p95'], db['mean'],
            db['p50'], db['p95'], dl['mean'], dl['p50'], dl['p95'],
            sum(v > 1e-9 for v in lead), sum(abs(v) <= 1e-9 for v in lead),
            sum(v < -1e-9 for v in lead),
            mean(bb) / mean(aa) if aa and mean(aa) > 0 else None, note]


def pair_runs(left, right):
    if left['method'] == METHODS[2] and right['method'] == METHODS[1]:
        audit = compare_natural_baseline(left['meta'], left['rows'], right['meta'], right['rows'])
        if audit['paired_identical_output_and_nfe_count'] != len(left['rows']):
            raise ValueError('OMTP/Natural output or NFE mismatch; review before export')
        warning = '; '.join(audit['warnings'])
    else:
        audit = compare_ar(left['meta'], left['rows'], right['meta'], right['rows'])
        warning = '; '.join(audit['warnings'])
    rr = {r['sample_id']: r for r in right['rows']}
    pairs = [(r, rr[r['sample_id']]) for r in left['rows'] if r['sample_id'] in rr]
    out = []
    def add(metric, target, values, note=''):
        out.append(paired_rows(left['dataset'], left['method'], right['method'], values,
                               target, metric, note + ' ' + warning))
    add('共同最终工具的首次识别时间', 'Final(A)∩Final(B)，逐工具去重',
        [(a['online'].get(t), b['online'].get(t)) for a, b in pairs for t in sorted(a['final'] & b['final'])],
        '同请求、同一工具名称配对；可能不是GT正确工具；差值=B-A，正值A更早。')
    for reference, label in [('final', '同一非空Final工具集合'), ('gt', '同一Ground Truth工具集合')]:
        chosen = [(a, b) for a, b in pairs if a[reference] and
                  (reference == 'gt' or a['final'] == b['final'])]
        add('完整集合覆盖时间', label,
            [(complete_time(a['online'], a[reference]), complete_time(b['online'], b[reference])) for a, b in chosen],
            '覆盖不等于精确匹配，额外错误触发另外报告；差值=B-A，正值A更早。')
    for k in range(1, 5):
        times = []
        for a, b in pairs:
            if len(a['gt']) < k:
                continue
            av = sorted(a['online'][t] for t in a['gt'] if t in a['online'])
            bv = sorted(b['online'][t] for t in b['gt'] if t in b['online'])
            times.append((av[k - 1] if len(av) >= k else None, bv[k - 1] if len(bv) >= k else None))
        add(f'第{k}个正确工具识别时间', 'Ground Truth；按识别时间排序，非调用位置', times,
            '同一请求的第k个正确类型，两方法第k个工具的名称可能不同。')
    return out


def build(output, results_root):
    if output.exists():
        raise FileExistsError(f'refusing to overwrite {output}')
    runs = []
    for dataset, (domain, paths) in DATASETS.items():
        samples = sample_catalog(domain)
        for method in METHODS:
            runs.append(analyze(dataset, domain, method, results_root / paths[method] / 'samples.jsonl', samples))
        print(f'Analyzed {dataset}: 3 methods × {len(samples)} samples', flush=True)
    book = Workbook()
    book.remove(book.active)
    guides = [
        ['范围', '4个数据集 × 3种方法 × 200条请求 = 2400条运行记录；独立query为800条。只汇总本地现有正式结果，不混入Smoke或历史配置。'],
        ['方法', 'AR Natural = Meta-Llama-3-8B-Instruct + HF原生KV Cache；LLaDA Natural = LLaDA-1.5-8B + Natural-only；OMTP + Fallback = 同一LLaDA + optimized OMTP/prefix/natural。'],
        ['先看哪里', '总览看方法概况；公平配对时间/三方法共同工具时间看横向时间比较；准确率_自定义区分Final与GT；官方指标状态注明尚未评测的官方项目。'],
        ['时间零点', '生成调用前；包含模型处理Prompt、生成和在线检测；不包含模型加载、Prompt渲染/分词、离线Repair/统计及日志写盘。单位秒。'],
        ['时间样本', '逐工具以query×工具名称为单位，重复调用只算第一次；第1–4个正确工具按检测时间排序，不是数组位置。每项均列有效数与失败数，失败为留空，绝不填0。'],
        ['均值/分位数', '均值为算术平均；P50/P95为排序后线性插值，索引=(n-1)×p；负提前量保留。配对提前量的分位数直接对逐对差值计算，不用两组分位数相减。'],
        ['真实净提前量', '独立Natural-only运行时间减OMTP-on实际触发时间，包含OMTP观测开销；不是同轮Natural时间减OMTP时间，也不是直接从提前量减全部Observer耗时。'],
        ['准确率口径', 'Final Raw Text是当前extract_partial_names规则识别到的完整JSON-like name字段，并限制在合法候选目录内；此口径可分析残缺JSON但不代表调用可执行，也不覆盖未知非法工具名。'],
        ['任务正确性', '在线集合相对GT为自定义工具类型指标；最终生成质量单独解析严格调用、重复次数、参数。它们不是同一指标，OMTP的GT F1改变不代表模型最终输出改变。'],
        ['官方指标缺口', '12个日志均未记录官方evaluator成绩，环境无bfcl_eval。官方指标值明确为未运行，不能拿工具Set EM、本地AST或语义边F1代替。'],
        ['本地兼容评测', '生成质量表重新调用现有parse_tool_calls / json-repair / bfcl_compatible_ast_match / semantic dependency checker，属于离线兼容评测，不改原始输出、事件、时间。'],
        ['JSON Repair', f'使用已安装json-repair {version("json-repair")}；仅对不合法JSON运行一次；使用output_before_eos（若有）否则raw_output。Raw全文JSON合法率另列。'],
        ['官方可比性', 'TaskBench只选Chain/DAG、至少2个调用且至少2种工具的固定200条；输出为自定义flat_name_v1。不是TaskBench完整官方评测设置；转换与子集需在论文中披露。'],
        ['调用位置/次数缺口', '当前tool_first_v1仅保存每工具类型首次事件，不能恢复OMTP occurrence次数与逐位置预测准确率；对应栏标不可评估。生成最终列表的次数/位置只作另外的离线诊断。'],
        ['顺序语义', 'BFCL parallel_multiple的官方调用匹配不要求数组顺序；DAG也可能有多个合法拓扑顺序。严格逐位置比较不能代替官方任务正确性。'],
        ['停止条件', 'LLaDA固定128/128/32；AR不设实验输出预算，生成到EOS或物理上下文上限。AR context_limit不是自然结束；所有失败仍保留在全样本分母。'],
        ['参数差异', 'TaskBench OMTP strict_anchor=true、append_natural=true、score>=-5、margin>=0.1；BFCL正式结果关闭这些门控。工作簿忠实记录，不声称四领域完全同一门控配置。'],
        ['Parity', '每轮仅首条做Observer On/Off检查，未检查的样本不能当作全量Parity通过；另核对LLaDA两轮200条Raw output与NFE一致性。'],
        ['时间限制', '均为真实历史墙钟数据，未独立监控共享GPU/系统负载，不把所有跨轮时间差归因Observer；AR与dLLM NFE不代表等量计算。'],
        ['不作的宣称', '检测时间比不是工具执行端到端加速比；实验未连接真实工具启动，也未测量端到端预热收益。'],
        ['可复现', f'/home/yzx/miniconda3/envs/llada/bin/python {ROOT / "export_paper_workbook.py"} --output <新的.xlsx路径>；输入JSONL SHA256和命令见结果溯源。'],
        ['生成时间', datetime.now().astimezone().isoformat()],
    ]
    sheet(book, '阅读说明', ['项目', '说明'], guides)
    overview, configs, datasets, times, accuracies, generation, official, audit_rows = [], [], [], [], [], [], [], []
    sample_rows, tool_rows, groups = [], [], []
    for run in runs:
        ds, method, rows, meta = run['dataset'], run['method'], run['rows'], run['meta']
        a = meta['args']; n = len(rows)
        qf, qg = set_metrics(run, 'final'), set_metrics(run, 'gt')
        gen, obs = (distribution([r.get(k) for r in rows]) for k in ['generation_time_s', 'observer_time_s'])
        tt = distribution([r['online'].get(t) for r in rows for t in r['final']])
        gen_rows = generation_rows(run)
        overview.append([ds, method, n, gen['mean'], gen['p50'], gen['p95'],
                         tt['count'], tt['mean'], tt['p50'], tt['p95'],
                         qf['precision'], qf['recall'], qf['f1'], qf['exact'], qf['wrong_rate'],
                         qg['precision'], qg['recall'], qg['f1'], qg['exact'],
                         gen_rows[0][4], gen_rows[0][11], obs['mean'],
                         NA_OFFICIAL if run['domain'] == 'bfcl' else '不适用',
                         NA_OFFICIAL if run['domain'] != 'bfcl' else '不适用'])
        times.extend(timing_rows(run)); accuracies.extend(quality_rows(run)); generation.extend(gen_rows)
        algorithm = meta.get('algorithm', meta.get('online_method', {}))
        configs.append([ds, method, a.get('model_path_resolved', a.get('model_path')),
                        'HF AR / KV Cache' if method == METHODS[0] else 'Fast-dLLM v1 / No Cache',
                        a.get('gen_length'), a.get('block_length'), a.get('steps'),
                        a.get('temperature'), a.get('seed'), a.get('dtype'),
                        a.get('prompt_profile', 'original'), a.get('output_schema', 'standard'),
                        'EOS / context_limit（无输出预算）' if method == METHODS[0] else '固定128画布',
                        json.dumps(algorithm, ensure_ascii=False),
                        sum(bool(r.get('error')) for r in rows),
                        sum(r.get('parity', {}).get('checked', False) for r in rows),
                        sum(r.get('parity', {}).get('output_identical') is False or
                            r.get('parity', {}).get('nfe_identical') is False for r in rows),
                        mean(r['forward_count'] for r in rows if r.get('forward_count') is not None),
                        min(r['forward_count'] for r in rows if r.get('forward_count') is not None),
                        max(r['forward_count'] for r in rows if r.get('forward_count') is not None),
                        sum(r.get('stop_reason') == 'context_limit' for r in rows),
                        sum(bool(r.get('eos_detected')) for r in rows)])
        digest = hashlib.sha256(run['path'].read_bytes()).hexdigest()
        audit_rows.append([ds, method, str(run['path'].resolve()), digest,
                           str(meta.get('created_at')), meta.get('gpu'), meta.get('torch'), meta.get('cuda'),
                           meta.get('fastdllm_commit'), json.dumps(meta.get('code_sha256', {}), ensure_ascii=False),
                           json.dumps(meta.get('command'), ensure_ascii=False),
                           json.dumps(meta.get('timing_semantics'), ensure_ascii=False)])
        metrics = ['AST Accuracy (parallel_multiple)'] if run['domain'] == 'bfcl' else ['n-F1', 'e-F1', 't-F1', 'v-F1']
        for metric in metrics:
            for mode in ['Original', 'JSON Repair（额外后处理条件）']:
                official.append([ds, method, mode, metric, None, NA_OFFICIAL,
                                 '无官方evaluator输出；本地兼容分数见生成质量，禁止替代官方值',
                                 OFFICIAL_LINKS['bfcl' if run['domain'] == 'bfcl' else 'taskbench']])
        if method == METHODS[0]:
            refs = list(run['samples'].values())
            datasets.append([ds, n, sum(getattr(s, 'graph_type', None) == 'chain' for s in refs),
                             sum(getattr(s, 'graph_type', None) == 'dag' for s in refs),
                             min(len(s.candidate_names) for s in refs), max(len(s.candidate_names) for s in refs),
                             min(len(s.ground_truth_calls) for s in refs), max(len(s.ground_truth_calls) for s in refs),
                             sum(len(s.ground_truth_calls) for s in refs),
                             sum(len(s.ground_truth_tool_types) for s in refs),
                             'BFCL非Live完整parallel_multiple' if run['domain'] == 'bfcl' else
                             '固定formal200；Chain/DAG各100；>=2调用且>=2工具类型；与Smoke20互斥',
                             json.dumps({k: v for k, v in meta.items() if k.endswith('_sha256') and k != 'code_sha256'}, ensure_ascii=False)])
        if run['domain'] != 'bfcl':
            for graph in ['chain', 'dag']:
                subset = [r for r in rows if r.get('graph_type') == graph]
                for reference, label in [('final', 'Final Raw Text'), ('gt', 'Ground Truth')]:
                    q = quality([set(r['online']) for r in subset], [r[reference] for r in subset])
                    dt = distribution([r['online'].get(t) for r in subset for t in r[reference]])
                    groups.append([ds, method, graph, label, len(subset), q['precision'], q['recall'],
                                   q['f1'], q['exact'], q['covered'], q['wrong_per_query'],
                                   dt['count'], dt['mean'], dt['p50'], dt['p95']])
        for r in rows:
            q1 = quality([set(r['online'])], [r['final']]); q2 = quality([set(r['online'])], [r['gt']])
            sample_rows.append([ds, method, r['sample_id'], r.get('graph_type'),
                                len(r['gt_names']), len(r['gt']), len(r['final']), len(r['online']),
                                q1['fp'], q1['fn'], q1['exact'], q2['fp'], q2['fn'], q2['exact'],
                                r.get('generation_time_s'), r.get('observer_time_s'), r.get('forward_count'),
                                complete_time(r['online'], r['final']), complete_time(r['online'], r['gt']),
                                r['strict'].json_valid, r['strict'].valid, r['repaired'].json_valid,
                                r['repaired'].valid, r['strict_ast'], r['repaired_ast'],
                                r.get('stop_reason'), r.get('eos_detected'), r.get('error'),
                                ', '.join(sorted(r['gt'])), ', '.join(sorted(r['final'])),
                                ', '.join(sorted(r['online'])), r['strict'].error])
            first = {e['tool']: e for e in r['first_tool_events']}
            for tool in sorted(set(r['online']) | r['final'] | r['gt']):
                e = first.get(tool, {})
                tool_rows.append([ds, method, r['sample_id'], tool, tool in r['gt'], tool in r['final'],
                                  tool in r['online'], e.get('source'), r['online'].get(tool),
                                  r['natural'].get(tool), e.get('forward'),
                                  r['natural'][tool] - r['online'][tool] if tool in r['natural'] and tool in r['online'] else None,
                                  '同轮Natural时间包含OMTP开销，不是独立基线净收益'])
    sheet(book, '总览', ['数据集', '方法', '请求样本数', '生成时间 Mean(s)', '生成时间 P50(s)', '生成时间 P95(s)',
          'Final逐工具有效数', 'Final逐工具时间 Mean(s)', 'Final逐工具时间 P50(s)', 'Final逐工具时间 P95(s)',
          'Final Trigger Precision', 'Final Trigger Recall', 'Final Trigger F1', 'Final Tool Set Exact Match',
          'Final错误触发率', 'GT在线 Precision', 'GT在线 Recall', 'GT在线 F1', 'GT在线 Tool Set Exact Match',
          'Raw全文JSON合法率', '严格生成GT工具集合准确率', 'Observer Mean(s)', '官方BFCL AST', '官方TaskBench n-F1'], overview)
    sheet(book, '数据集', ['数据集', '请求样本数', 'Chain样本数', 'DAG样本数', '候选工具数最小值', '候选工具数最大值',
          'GT调用数最小值', 'GT调用数最大值', 'GT调用总次数', 'GT去重工具总数', '筛选说明', '数据来源哈希'], datasets)
    sheet(book, '方法配置与检查', ['数据集', '方法', '模型路径', '推理方式', 'gen_length', 'block_length', 'steps',
          'temperature', 'seed', 'dtype', 'prompt_profile', 'output_schema', '停止条件', '算法配置说明',
          '运行失败样本数', 'Parity已检查样本数', 'Parity失败样本数', 'NFE Mean', 'NFE Min', 'NFE Max',
          'AR上下文上限失败样本数', '检测到EOS样本数'], configs)
    sheet(book, '时间指标', ['数据集', '方法', '时间指标', '参考目标', '符合条件数量', '有效样本数', '缺失或失败数量',
          '未成功比例', 'Mean(s)', 'P50(s)', 'P95(s)', '口径说明'], times)
    sheet(book, '准确率_自定义', ['数据集', '方法', '正确性参考', '请求样本数', 'TP工具数', 'FP工具数', 'FN工具数',
          'Precision micro', 'Recall micro', 'F1 micro', '集合正确请求数', 'Tool Set Exact Match',
          '完整覆盖请求数', '非空目标请求数', '完整集合覆盖比例', '错误触发率', '每请求错误工具种数',
          '错误触发请求数', '错误触发请求比例', '每请求遗漏工具种数', '在线Occurrence F1状态', '在线逐位置准确率状态'], accuracies)
    sheet(book, '生成质量_兼容非官方', ['数据集', '方法', '解析条件', '请求样本数', 'Raw全文JSON合法率',
          '评测文本JSON合法率', '调用结构合法率', '完整合法输出率', '生成工具 Precision', '生成工具 Recall',
          '生成工具 F1', '严格工具集合GT准确率', 'Raw Text工具集合GT准确率', 'Occurrence Precision', 'Occurrence Recall',
          'Occurrence F1', '调用总次数匹配率', '每工具调用次数匹配率', '调用位置准确率_仅诊断',
          '本地AST兼容准确率_非官方', '本地语义边 F1_非官方', '官方评测状态', '评测说明'], generation)
    sheet(book, '官方指标状态', ['数据集', '方法', '后处理条件', '官方指标名称', '官方指标值', '官方评测状态',
          '缺失原因', '官方参考链接'], official)
    pair_data, common_data = [], []
    for dataset in DATASETS:
        rr = {r['method']: r for r in runs if r['dataset'] == dataset}
        pair_data.extend(pair_runs(rr[METHODS[2]], rr[METHODS[1]]))
        pair_data.extend(pair_runs(rr[METHODS[1]], rr[METHODS[0]]))
        pair_data.extend(pair_runs(rr[METHODS[2]], rr[METHODS[0]]))
        indexed = {m: {r['sample_id']: r for r in rr[m]['rows']} for m in METHODS}
        for correct_only in [False, True]:
            values = {m: [] for m in METHODS}; eligible = 0
            for sid in indexed[METHODS[0]]:
                allrows = [indexed[m][sid] for m in METHODS]
                target = set.intersection(*(r['final'] for r in allrows))
                if correct_only:
                    target &= allrows[0]['gt']
                eligible += len(target)
                for tool in sorted(target):
                    if all(tool in r['online'] for r in allrows):
                        for method, row in zip(METHODS, allrows):
                            values[method].append(row['online'][tool])
            for method in METHODS:
                d = distribution(values[method])
                common_data.append([dataset, method, '三方法共同生成且GT正确的工具' if correct_only else '三方法共同生成的工具（不保证GT正确）',
                                    eligible, d['count'], eligible - d['count'], d['mean'], d['p50'], d['p95'],
                                    '同请求×同工具名称；仅三方均成功的条件子集，不能当全样本任务完成时间'])
    sheet(book, '公平配对时间', ['数据集', '方法A', '方法B', '配对目标', '时间指标', '符合条件数量', '双方成功数量',
          '仅A成功数量', '仅B成功数量', '双方失败数量', 'A Mean(s)', 'A P50(s)', 'A P95(s)',
          'B Mean(s)', 'B P50(s)', 'B P95(s)', '提前量B减A Mean(s)', '提前量B减A P50(s)', '提前量B减A P95(s)',
          'A更早数量', '相同时间数量', 'A更晚数量', '识别时间比B除A', '配对说明及警告'], pair_data)
    sheet(book, '三方法共同工具时间', ['数据集', '方法', '配对目标', '符合条件工具数', '三方成功工具数', '缺失工具数',
          'Mean(s)', 'P50(s)', 'P95(s)', '口径说明'], common_data)
    sheet(book, 'Chain_DAG分组', ['数据集', '方法', '分组', '正确性参考', '请求样本数', 'Precision micro', 'Recall micro',
          'F1 micro', 'Tool Set Exact Match', '完整集合覆盖比例', '每请求错误工具种数', '有效时间工具数',
          '逐工具时间 Mean(s)', '逐工具时间 P50(s)', '逐工具时间 P95(s)'], groups)
    sheet(book, '逐样本核对', ['数据集', '方法', 'sample_id', 'graph_type', 'GT调用数', 'GT工具类型数', 'Final工具类型数',
          '触发工具类型数', 'Final FP数', 'Final FN数', 'Final集合完全一致', 'GT FP数', 'GT FN数', 'GT集合完全一致',
          '生成时间(s)', 'Observer时间(s)', 'NFE', 'Final完整覆盖时间(s)', 'GT完整覆盖时间(s)',
          '原始JSON合法', '原始完整调用合法', 'Repair JSON合法', 'Repair完整调用合法', '原始本地AST兼容',
          'Repair本地AST兼容', 'AR停止原因', '检测到EOS', '运行错误', 'GT工具名称', 'Final工具名称', '触发工具名称',
          '原始解析错误'], sample_rows)
    sheet(book, '逐工具核对', ['数据集', '方法', 'sample_id', '工具名称', 'GT包含', 'Final包含', '已触发', '最早触发来源',
          '首次在线时间(s)', '同轮Natural时间(s)', '触发Forward', '同轮提前量(s)', '口径说明'], tool_rows)
    sheet(book, '结果溯源', ['数据集', '方法', '输入文件路径', '输入SHA256哈希', '运行创建时间', 'GPU', 'PyTorch', 'CUDA',
          'FastDLLM版本', '代码SHA256哈希', '实际运行命令', '原始计时语义说明'], audit_rows)
    definitions = [
        ['Trigger Precision', 'Σ|P∩F| / Σ|P|；P为不可撤回的首次触发集合，F为同一模型最终Raw Text候选工具集合。', '自定义，排除模型任务选择错误影响'],
        ['Trigger Recall', 'Σ|P∩F| / Σ|F|。全量Fallback可能使其接近100%，不代表提前识别覆盖率。', '自定义'],
        ['F1 micro', '2TP/(2TP+FP+FN)，在全部query×工具类型上聚合；不是逐请求F1的平均。', '自定义'],
        ['Tool Set Exact Match', '所有请求中P=目标集合的比例；不忽略额外工具。', '自定义，GT与Final分开'],
        ['完整集合覆盖比例', '非空目标请求中目标集合⊆P的比例；允许额外工具，不等于Exact Match。空目标单独排除并报告分母。', '自定义'],
        ['错误触发率', 'ΣFP/(ΣTP+ΣFP)；按每query每工具首次事件计数，不是每个Forward或轨迹commit计数。', '自定义'],
        ['错误触发请求比例', '至少一次错误工具首次触发的请求数/全部请求数。', '自定义'],
        ['逐工具时间', '每query×目标工具名称的首次有效事件时间；OMTP/prefix/natural谁先采用谁，不依赖GT在线挑选。', '日志内在线时刻，正确性离线标注'],
        ['完整集合时间', '目标每种工具首次事件时间的最大值；若漏掉任一工具则缺失。', '离线oracle目标，不是在线已知集合完成信号'],
        ['配对提前量', '同一目标下T_B-T_A；正值A更早。双成功对直接算差再求分位数；其余成功组合全部计数。', '秒，包含真实各自Observer开销'],
        ['Occurrence F1', '生成最终调用列表按工具名称Counter交集计TP，重复调用保留。分母为全部可解析预测和GT调用。', '仅最终生成；不等于在线occurrence'],
        ['调用总次数匹配率', '完整合法输出且最终调用列表长度=GT调用列表长度的请求比例；不能保证各工具次数相同。', '生成质量诊断'],
        ['每工具调用次数匹配率', '完整合法输出且预测名称Counter=GT名称Counter的请求比例。', '生成质量诊断'],
        ['调用位置准确率', '按列表索引对应，名称相等的位置数/GT总调用数；缺失位置算错，多余调用另看次数/P。', '仅诊断，不适合作BFCL并行任务主正确率'],
        ['严格工具集合准确率', '完整合法调用输出且生成名称集合=GT集合的请求比例；格式失败算失败。', '自定义；不是BFCL AST'],
        ['本地AST兼容准确率', '现有bfcl_compatible_ast_match，无序、重复次数敏感、名称及参数/类型/允许值匹配；要求输出完整合法。', '本地兼容，非官方；TaskBench参数为适配后字面匹配'],
        ['本地语义边F1', '生成参数中<node-N>解析成工具名依赖边，与TaskBench edges比较；去重边micro F1。', '本地兼容，未验证等同官方e-F1'],
        ['官方BFCL AST', '应使用与数据版本一致的BFCL官方evaluator评测确定性转换后的调用。', '未运行，不填估计分数'],
        ['官方TaskBench n/e/t/v-F1', '工具节点/依赖边/参数名称/参数值F1；须使用官方evaluator及明确输出适配。', '未运行；去重集合F1不能自动冒充n-F1'],
        ['未触发/失败时间', '空白表示没有成功观测到，不是0秒；统计均值分位数必须同时看有效数和失败数。', 'NA'],
    ]
    sheet(book, '指标定义', ['指标', '定义说明', '适用范围说明'], definitions)
    # Make the workbook self-navigating without adding any external artifacts.
    guide = book['阅读说明']
    for name in book.sheetnames[1:]:
        guide.append([f'打开：{name}', name])
        guide.cell(guide.max_row, 2).hyperlink = f"#'{name}'!A1"
        guide.cell(guide.max_row, 2).style = 'Hyperlink'
    for ws in [book['总览'], book['准确率_自定义']]:
        ws.cell(1, 1).comment = Comment('自定义集合指标不是官方benchmark成绩。请先阅读“阅读说明”。', 'Codex')
    output.parent.mkdir(parents=True, exist_ok=True)
    book.save(output)
    check = load_workbook(output, read_only=True, data_only=True)
    assert check['总览'].max_row == 13
    assert check['逐样本核对'].max_row == 2401
    assert all(row[4] is None and row[5] == NA_OFFICIAL for row in check['官方指标状态'].iter_rows(min_row=2, values_only=True))
    counts = {ws.title: ws.max_row - 1 for ws in check}
    check.close()
    print(json.dumps({'output': str(output.resolve()), 'sheets_data_rows': counts,
                      'bytes': output.stat().st_size, 'official_scores': NA_OFFICIAL}, ensure_ascii=False, indent=2))


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--results-root', type=Path, default=ROOT / 'results')
    parser.add_argument('--output', type=Path, required=True)
    args = parser.parse_args()
    build(args.output.resolve(), args.results_root.resolve())


if __name__ == '__main__':
    main()
